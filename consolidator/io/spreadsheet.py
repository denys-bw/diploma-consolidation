from __future__ import annotations

import io
from datetime import date, datetime, time
from typing import Any

import pandas as pd
from openpyxl import load_workbook

from consolidator.io.formats import RawFile, ReadOptions, TableFormat
from consolidator.io.names import stem_of, unique_column_names
from consolidator.io.source import Source


class SpreadsheetFormat(TableFormat):
    suffixes = (".xlsx", ".xlsm")

    def read(self, file: RawFile, options: ReadOptions) -> list[Source]:
        sheets = parse_xlsx(file.data)
        return [
            Source(title if len(sheets) > 1 else stem_of(file.name), frame, file.origin, "XLSX", {"sheet": title})
            for title, frame in sheets
        ]


def parse_xlsx(data: bytes) -> list[tuple[str, pd.DataFrame]]:
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        sheets = [(sheet.title, _sheet_frame(sheet.iter_rows(values_only=True))) for sheet in workbook.worksheets]
    finally:
        workbook.close()
    return [(title, frame) for title, frame in sheets if frame is not None]


def cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time(0) else value.isoformat(sep=" ")
    if isinstance(value, (date, time)):
        return value.isoformat()
    return str(value)


def _sheet_frame(rows: Any) -> pd.DataFrame | None:
    values = [[cell_text(cell) for cell in row] for row in rows]
    values = [row for row in values if any(cell.strip() for cell in row)]
    if not values:
        return None
    width = max(len(row) for row in values)
    while width and not any(len(row) >= width and row[width - 1] for row in values):
        width -= 1
    rows_fitted = [(row + [""] * width)[:width] for row in values]
    header, body = rows_fitted[0], rows_fitted[1:]
    return pd.DataFrame(body, columns=unique_column_names(header), dtype=object)
