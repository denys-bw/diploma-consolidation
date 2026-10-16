import json
import sqlite3
from contextlib import closing

import pandas as pd
import pytest
from openpyxl import load_workbook

from consolidator.imputation import MeanImputer
from consolidator.io import Source
from consolidator.linkage import ClericalReview, FullComparison, LinkageSettings
from consolidator.pipeline import Pipeline, PipelineSettings
from consolidator.profiling import ColumnType
from consolidator.reporting import (
    export_csv,
    export_excel,
    export_json,
    export_sqlite,
    profile_table,
    quality_table,
    report_tables,
    summarize,
)
from consolidator.standardization import StandardizationSettings


def guides() -> list[Source]:
    fodors = pd.DataFrame(
        {
            "name": [
                "arnie morton's of chicago",
                "art's delicatessen",
                "hotel bel-air",
                "cafe bizou",
                "campanile",
                "chinois on main",
            ],
            "addr": [
                "435 s. la cienega blv.",
                "12224 ventura blvd.",
                "701 stone canyon rd.",
                "14016 ventura blvd.",
                "624 s. la brea ave.",
                "2709 main st.",
            ],
            "city": ["los angeles", "studio city", "bel air", "sherman oaks", "los angeles", "santa monica"],
            "phone": ["310/246-1501", "818/762-1221", "310/472-1211", "818/788-3536", "213/938-1447", ""],
        }
    )
    zagats = pd.DataFrame(
        {
            "name": ["arnie morton's of chicago", "art's deli", "bel-air hotel", "cafe bizou", "patina"],
            "address": [
                "435 s. la cienega blvd.",
                "12224 ventura blvd.",
                "701 stone canyon rd.",
                "14016 ventura blvd.",
                "5955 melrose ave.",
            ],
            "city": ["los angeles", "studio city", "bel air", "sherman oaks", "los angeles"],
            "phone": ["310-246-1501", "818-762-1221", "310-472-1211", "818-788-3536", "213-467-1108"],
        }
    )
    return [Source("fodors", fodors, "fodors.csv", "CSV"), Source("zagats", zagats, "zagats.json", "JSON")]


SETTINGS = PipelineSettings(
    standardization=StandardizationSettings(phone_format="999-999-9999"),
    linkage=LinkageSettings(blocker=FullComparison()),
)


class StageLog:
    def __init__(self):
        self.stages = []

    def stage_started(self, stage):
        self.stages.append(stage)


@pytest.fixture(scope="module")
def result():
    return Pipeline(SETTINGS).run(guides())


def test_pipeline_matches_columns_and_links_records(result):
    assert result.attributes == ["name", "addr", "city", "phone"]
    assert result.mapping.attribute("addr").column_in("zagats") == "address"
    assert result.mapping.types["phone"] is ColumnType.PHONE
    assert result.fusion.n_entities == 7
    assert result.consolidated.loc[0, "sources"] == "fodors, zagats"
    phones = result.consolidated["phone"]
    assert phones[phones != ""].str.match(r"^\d{3}-\d{3}-\d{4}$").all()
    assert set(result.timings) == {
        "profiling",
        "schema",
        "standardization",
        "linkage",
        "fusion",
        "imputation",
        "quality",
    }


def test_pipeline_measures_quality_before_and_after(result):
    assert result.quality_before["uniqueness"] == pytest.approx(7 / 11)
    assert result.quality_before["consistency"] == pytest.approx(5 / 10)
    assert result.quality_before["completeness"] == pytest.approx(43 / 44)
    assert result.quality_after["uniqueness"] == 1.0
    assert result.quality_after["consistency"] == 1.0
    assert result.quality_after["completeness"] == pytest.approx(27 / 28)
    assert result.index_after > result.index_before


def test_pipeline_reports_each_stage_before_it_starts():
    log = StageLog()
    Pipeline(SETTINGS, observers=[log]).run(guides())
    assert log.stages == ["profiling", "schema", "standardization", "linkage", "fusion", "imputation", "quality"]


def test_review_decisions_reach_the_linkage_stage():
    review = ClericalReview.of([(("fodors", "0"), ("zagats", "0"), False)])
    result = Pipeline(SETTINGS).run(guides(), review=review)
    assert result.fusion.n_entities == 8
    assert result.consolidated.loc[0, "sources"] == "fodors"


def test_record_values_follow_the_entities(result):
    values = result.record_values()
    assert len(values) == 11
    rows = values.set_index(["source", "key"])
    assert rows.loc[("zagats", "1"), "name"] == rows.loc[("fodors", "1"), "name"]
    assert rows.loc[("fodors", "5"), "phone"] == ""


def test_same_columns_use_one_profile_for_types():
    frame = pd.DataFrame({"src": ["a", "a", "b", "b"], "dep": ["7:10 a.m.", "8:00 a.m.", "7:10 a.m.", ""]})
    sources = Source("flights", frame, "flights.csv", "CSV").split_by("src")
    mapping = Pipeline().match(sources)
    assert mapping.names == ["dep"]
    assert mapping.types["dep"] is ColumnType.TIME


def test_reports_are_built_and_exported(result, tmp_path):
    tables = report_tables(result)
    assert list(tables)[:3] == ["Консолідовані дані", "Походження значень", "Якість"]
    quality = quality_table(result).set_index("Критерій")
    assert quality.loc["Інтегральний індекс якості", "Після обробки"] == pytest.approx(result.index_after, abs=1e-4)
    assert not profile_table(result.profiles["fodors"]).empty
    summary = summarize(result)
    assert summary.record_count == 11
    assert summary.entity_count == 7
    assert summary.merged_records == 4
    workbook = load_workbook(export_excel(result, tmp_path / "report.xlsx"), read_only=True)
    assert workbook.sheetnames == list(tables)
    assert [path.name for path in export_csv(result, tmp_path / "csv")] == [
        "consolidated.csv",
        "provenance.csv",
        "quality.csv",
    ]
    payload = json.loads(export_json(result, tmp_path / "report.json").read_text(encoding="utf-8"))
    assert payload["summary"]["entity_count"] == 7


def test_result_is_saved_to_database(result, tmp_path):
    path = export_sqlite(result, tmp_path / "result.sqlite")
    with closing(sqlite3.connect(path)) as connection:
        tables = [row[0] for row in connection.execute("SELECT name FROM sqlite_master WHERE type = 'table'")]
        merged = connection.execute(
            'SELECT e."Сутність", e.name, COUNT(*) FROM entities e '
            'JOIN source_records r ON r."Сутність" = e."Сутність" GROUP BY e."Сутність" HAVING COUNT(*) > 1'
        ).fetchall()
        missing = connection.execute("SELECT COUNT(*) FROM entities WHERE phone IS NULL").fetchone()[0]
        violations = connection.execute("PRAGMA foreign_key_check").fetchall()
    assert tables == ["entities", "provenance", "source_records", "quality"]
    assert len(merged) == 4
    assert all(count == 2 for _, _, count in merged)
    assert missing == int((result.consolidated["phone"] == "").sum())
    assert violations == []


def test_imputation_can_be_switched_off():
    settings = PipelineSettings(linkage=SETTINGS.linkage, imputer=None)
    result = Pipeline(settings).run(guides())
    assert result.imputation is None
    settings = PipelineSettings(linkage=SETTINGS.linkage, imputer=MeanImputer())
    assert Pipeline(settings).run(guides()).imputation.count == 0


def test_duplicate_search_can_be_switched_off():
    result = Pipeline(PipelineSettings(linkage=None)).run(guides())
    assert result.linkage is None
    assert result.fusion.n_entities == 11
    assert result.quality_after["uniqueness"] == 1.0
    assert summarize(result).candidate_pairs == 0
    assert "Пошук дублікатів" in report_tables(result)["Зв'язування"]["Показник"].tolist()
