from __future__ import annotations

import json
import re
import sqlite3
import sys
import zipfile
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
RAW_FILE = ROOT / "data" / "raw" / "consolidation" / "chicago" / "chicago_early_childhood.csv"
OUTPUT = ROOT / "data" / "prepared" / "chicago_sources"
DATABASE = "state_registries.sqlite"
FIXED_TIME = datetime(2026, 10, 8)
PLAN = {
    "CPS_Early_Childhood_Portal_scrape.csv": "json",
    "ECE Chicago Find a School scrape.csv": "json",
    "chapin_dfss_providers_2011_070212.csv": "xlsx",
    "DFSS_AgencySiteLies_2012.csv": "xlsx",
    "purple_binder_early_childhood.csv": "csv",
    "NAEYC_accreditation.csv": "csv",
    "IDHS_child_care_provider_list.csv": "sqlite",
    "Ounce of Prevention - Site Locations.csv": "sqlite",
    "chapin_ounce_providers_2011_071112.csv": "sqlite",
}


def main() -> None:
    raw = pd.read_csv(RAW_FILE, dtype=str, keep_default_na=False)
    parts = split_sources(raw)
    if set(parts) != set(PLAN):
        sys.exit(f"Джерела в наборі не збігаються з планом: {sorted(set(parts) ^ set(PLAN))}")
    OUTPUT.mkdir(parents=True, exist_ok=True)
    (OUTPUT / DATABASE).unlink(missing_ok=True)
    rows = []
    for source, file_format in PLAN.items():
        frame = parts[source]
        path, restored = WRITERS[file_format](frame, Path(source).stem)
        if restored.shape != frame.shape or list(restored.columns) != list(frame.columns):
            sys.exit(f"Після запису {path.name} змінилася структура таблиці")
        if not (restored.to_numpy() == frame.to_numpy()).all():
            sys.exit(f"Після запису {path.name} значення відрізняються від оригіналу")
        rows.append({"джерело": Path(source).stem, "файл": path.name, "записів": len(frame), "колонок": frame.shape[1]})
    print(pd.DataFrame(rows).to_string(index=False))
    print(f"Усього записів: {sum(row['записів'] for row in rows)} (в оригіналі {len(raw)})")


def split_sources(raw: pd.DataFrame) -> dict[str, pd.DataFrame]:
    parts = {}
    for source, group in raw.groupby("Source", sort=False):
        group = group.drop(columns=["Source", "True Id"])
        filled = [column for column in group.columns if column == "Id" or (group[column] != "").any()]
        parts[source] = group[filled].reset_index(drop=True)
    return parts


def write_json(frame: pd.DataFrame, stem: str) -> tuple[Path, pd.DataFrame]:
    path = OUTPUT / f"{stem}.json"
    records = [{key: value if value != "" else None for key, value in row.items()} for row in frame.to_dict("records")]
    path.write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8")
    restored = pd.DataFrame(json.loads(path.read_text(encoding="utf-8"))).fillna("").astype(str)
    return path, restored


def write_xlsx(frame: pd.DataFrame, stem: str) -> tuple[Path, pd.DataFrame]:
    path = OUTPUT / f"{stem}.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = stem[:31]
    sheet.append(list(frame.columns))
    for row in frame.itertuples(index=False):
        sheet.append([value if value != "" else None for value in row])
    workbook.properties.created = FIXED_TIME
    workbook.save(path)
    _fix_archive_times(path)
    rows = list(load_workbook(path, read_only=True).active.iter_rows(values_only=True))
    return path, pd.DataFrame(rows[1:], columns=rows[0]).fillna("").astype(str)


def write_csv(frame: pd.DataFrame, stem: str) -> tuple[Path, pd.DataFrame]:
    path = OUTPUT / f"{stem}.csv"
    frame.to_csv(path, index=False, encoding="utf-8")
    return path, pd.read_csv(path, dtype=str, keep_default_na=False)


def write_sqlite(frame: pd.DataFrame, stem: str) -> tuple[Path, pd.DataFrame]:
    path = OUTPUT / DATABASE
    table = "_".join(stem.replace("-", " ").split())
    with sqlite3.connect(path) as connection:
        frame.replace("", None).to_sql(table, connection, index=False)
        restored = pd.read_sql(f'SELECT * FROM "{table}"', connection).fillna("").astype(str)
    return path, restored


def _fix_archive_times(path: Path) -> None:
    stamp = FIXED_TIME.strftime("%Y-%m-%dT%H:%M:%SZ").encode()
    with zipfile.ZipFile(path) as archive:
        entries = [(info.filename, archive.read(info.filename)) for info in archive.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in entries:
            if name == "docProps/core.xml":
                data = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*", rb"\g<1>" + stamp, data)
            info = zipfile.ZipInfo(name, date_time=FIXED_TIME.timetuple()[:6])
            info.compress_type = zipfile.ZIP_DEFLATED
            archive.writestr(info, data)


WRITERS = {"json": write_json, "xlsx": write_xlsx, "csv": write_csv, "sqlite": write_sqlite}


if __name__ == "__main__":
    main()
